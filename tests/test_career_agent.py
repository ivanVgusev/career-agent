import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from career_agent.adapters.hh import HHClient
from career_agent.adapters.job_boards import (
    ArbeitnowAdapter, JobicyAdapter, RemotiveAdapter,
)
from career_agent.config import SpeechConfig
from career_agent.domain.models import TranscriptionResult, Vacancy
from career_agent.domain.workflow import APPLICATION_WORKFLOW
from career_agent.exporters.excel import export_tracker
from career_agent.exporters.markdown import export_vacancy
from career_agent.migration import LegacyMigrator
from career_agent.repositories.sqlite import CareerRepository
from career_agent.services.jobs import JobService, MigrationRequired
from career_agent.services.transcription import TranscriptionService


class RepositoryTests(unittest.TestCase):
    def test_upsert_preserves_application_and_records_history(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = CareerRepository(Path(directory) / "career.db")
            vacancy = Vacancy(
                source="hh.ru", source_id="123", company="Example", title="LLM developer",
                url="https://hh.ru/vacancy/123", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00", score=80,
            )
            repository.upsert_vacancy(vacancy)
            repository.set_application_status("hh.ru", "123", "Подал", "Тест")
            vacancy.source_status = "Закрыта"
            repository.upsert_vacancy(vacancy)
            row = repository.list_vacancies()[0]
            self.assertEqual(row["application_status"], "Подал")
            self.assertEqual(row["source_status"], "Закрыта")
            with repository.connect() as database:
                history = database.execute("SELECT COUNT(*) FROM status_events").fetchone()[0]
            self.assertEqual(history, 1)

    def test_excel_is_derived_from_database(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = CareerRepository(Path(directory) / "career.db")
            repository.upsert_vacancy(Vacancy(
                source="hh.ru", source_id="7", company="Example", title="Python",
                url="https://hh.ru/vacancy/7", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00",
            ))
            output = export_tracker(repository, Path(directory) / "tracker.xlsx")
            sheet = load_workbook(output, read_only=True)["Вакансии"]
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet.cell(2, 1).value, "Example")

    def test_existing_excel_blocks_writes_before_migration(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tracker = root / "tracker.xlsx"
            tracker.touch()
            repository = CareerRepository(root / "career.db")
            with patch("career_agent.services.jobs.TRACKER_PATH", tracker):
                with self.assertRaises(MigrationRequired):
                    JobService(repository).ensure_writable()

    def test_markdown_keeps_summary_chance_and_full_description(self):
        with tempfile.TemporaryDirectory() as directory:
            vacancy = Vacancy(
                source="test", source_id="1", company="Example", title="AI Engineer",
                url="https://example.test/1", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00", score=85,
                description="Полное описание работодателя.",
                extra={"agent_summary": "Краткое резюме.", "hiring_chance": 70},
            )
            output = export_vacancy(vacancy, Path(directory))
            content = output.read_text(encoding="utf-8")
            self.assertIn("## Краткое резюме\n\nКраткое резюме.", content)
            self.assertIn("Шанс найма: 70/100", content)
            self.assertIn("## Полное описание\n\nПолное описание работодателя.", content)

    def test_markdown_source_copy_omits_analysis(self):
        with tempfile.TemporaryDirectory() as directory:
            vacancy = Vacancy(
                source="test", source_id="2", company="Example", title="AI Engineer",
                url="https://example.test/2", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00", score=91,
                direction="AI", description="Полная копия описания.",
                extra={
                    "agent_summary": "Этот текст не должен попасть в карточку.",
                    "hiring_chance": 80,
                    "source_copy_only": True,
                },
            )
            output = export_vacancy(vacancy, Path(directory))
            content = output.read_text(encoding="utf-8")
            self.assertNotIn("Краткое резюме", content)
            self.assertNotIn("Предварительная оценка", content)
            self.assertNotIn("Шанс найма", content)
            self.assertNotIn("Совпадения", content)
            self.assertIn("## Полное описание\n\nПолная копия описания.", content)

    def test_application_workflow_is_idempotent_and_tracks_state(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = CareerRepository(Path(directory) / "career.db")
            repository.upsert_vacancy(Vacancy(
                source="test", source_id="3", company="Example", title="AI Engineer",
                url="https://example.test/3", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00",
            ))
            repository.set_application_status("test", "3", "Рассматриваю")
            vacancies, inserted = repository.initialize_application_workflow(APPLICATION_WORKFLOW)
            self.assertEqual(vacancies, 1)
            self.assertEqual(inserted, len(APPLICATION_WORKFLOW))
            vacancies, inserted = repository.initialize_application_workflow(APPLICATION_WORKFLOW)
            self.assertEqual(vacancies, 1)
            self.assertEqual(inserted, 0)
            repository.set_application_workflow_step(
                "test", "3", "verify_source", "done", "Источник проверен",
            )
            rows = repository.list_application_workflow("test", "3")
            self.assertEqual(len(rows), len(APPLICATION_WORKFLOW))
            self.assertEqual(rows[0]["state"], "done")
            self.assertEqual(rows[0]["note"], "Источник проверен")

    def test_source_copy_mode_survives_refresh(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = CareerRepository(Path(directory) / "career.db")
            vacancy = Vacancy(
                source="test", source_id="4", company="Example", title="AI Engineer",
                url="https://example.test/4", source_status="Активна",
                checked_at="2026-10-13T12:00:00+03:00", score=85,
                direction="AI", description="Первая версия.",
                extra={"agent_summary": "Анализ", "hiring_chance": 70},
            )
            repository.upsert_vacancy(vacancy)
            repository.mark_source_copy("test", "4")
            vacancy.score = 95
            vacancy.direction = "AI/LLM"
            vacancy.description = "Обновлённая версия источника."
            vacancy.extra = {"agent_summary": "Новый анализ", "hiring_chance": 90}
            repository.upsert_vacancy(vacancy)
            row = repository.get_vacancy("test", "4")
            payload = json.loads(row["payload_json"])
            self.assertEqual(row["score"], 0)
            self.assertEqual(row["direction"], "")
            self.assertEqual(row["description"], "Обновлённая версия источника.")
            self.assertTrue(payload["extra"]["source_copy_only"])
            self.assertNotIn("agent_summary", payload["extra"])
            self.assertNotIn("hiring_chance", payload["extra"])


class ScoringTests(unittest.TestCase):
    def test_hh_normalization_is_remote_ai_fit(self):
        raw = {
            "id": "5", "name": "Junior LLM Python developer",
            "description": "<p>RAG, NLP, чат-бот и FastAPI</p>",
            "key_skills": [{"name": "Python"}],
            "schedule": {"id": "remote", "name": "Удаленная работа"},
            "experience": {"id": "between1And3", "name": "1–3 года"},
            "employment": {"name": "Полная занятость"},
            "employer": {"name": "Example"}, "area": {"name": "Россия"},
            "alternate_url": "https://hh.ru/vacancy/5", "archived": False,
        }
        vacancy = HHClient().normalize(raw)
        self.assertGreaterEqual(vacancy.score, 70)
        self.assertFalse(vacancy.assessment.rejected)

    def test_senior_title_is_rejected_even_with_strong_keywords(self):
        raw = {
            "id": "6", "name": "Sr. AI Engineer", "description": "<p>LLM RAG Python</p>",
            "schedule": {"id": "remote", "name": "Удаленная работа"},
            "experience": {"id": "between1And3", "name": "1–3 года"},
            "employer": {"name": "Example"}, "alternate_url": "https://hh.ru/vacancy/6",
            "archived": False,
        }
        vacancy = HHClient().normalize(raw)
        self.assertTrue(vacancy.assessment.rejected)

    def test_three_year_requirement_in_text_is_rejected(self):
        raw = {
            "id": "8", "name": "Python AI Engineer",
            "description": "<p>Python, LLM. Коммерческий опыт от 3 лет.</p>",
            "schedule": {"id": "remote", "name": "Удаленная работа"},
            "experience": {"id": "between1And3", "name": "1–3 года"},
            "employer": {"name": "Example"}, "alternate_url": "https://hh.ru/vacancy/8",
            "archived": False,
        }
        vacancy = HHClient().normalize(raw)
        self.assertTrue(vacancy.assessment.rejected)

    def test_plain_three_year_experience_requirement_is_rejected(self):
        raw = {
            "id": "10", "name": "Python AI Engineer",
            "description": "<p>Опыт работы 3 года. Python, LLM.</p>",
            "schedule": {"id": "remote", "name": "Удаленная работа"},
            "experience": {"id": "between1And3", "name": "1–3 года"},
            "employer": {"name": "Example"}, "alternate_url": "https://hh.ru/vacancy/10",
            "archived": False,
        }
        vacancy = HHClient().normalize(raw)
        self.assertTrue(vacancy.assessment.rejected)

    def test_middle_title_is_rejected(self):
        raw = {
            "id": "9", "name": "Middle Python Developer",
            "description": "<p>Python, LLM</p>",
            "schedule": {"id": "remote", "name": "Удаленная работа"},
            "experience": {"id": "between1And3", "name": "1–3 года"},
            "employer": {"name": "Example"}, "alternate_url": "https://hh.ru/vacancy/9",
            "archived": False,
        }
        vacancy = HHClient().normalize(raw)
        self.assertTrue(vacancy.assessment.rejected)


class ConnectorTests(unittest.TestCase):
    def test_remotive_normalization(self):
        client = Mock()
        client.get.return_value = {"jobs": [{
            "id": 10, "title": "Python LLM Engineer", "company_name": "Example",
            "description": "<p>RAG and FastAPI</p>", "url": "https://remotive.com/job/10",
            "candidate_required_location": "Worldwide", "job_type": "full_time",
            "salary": "USD 50000", "tags": ["Python", "AI"],
            "publication_date": "2026-10-13T10:00:00Z",
        }]}
        vacancy = RemotiveAdapter(client).search(["Python"], 10)[0]
        self.assertEqual(vacancy.source, "remotive")
        self.assertGreaterEqual(vacancy.score, 40)
        self.assertEqual(vacancy.area, "Worldwide")

    def test_jobicy_cursor_and_normalization(self):
        client = Mock()
        client.get.return_value = {"jobs": [{
            "id": 20, "jobTitle": "Data Analyst Python", "companyName": "Example",
            "jobDescription": "<p>SQL and analytics</p>", "url": "https://jobicy.com/jobs/20",
            "jobGeo": "EMEA", "jobType": ["full-time", "contract"],
            "jobIndustry": ["data-science"],
            "pubDate": "2026-10-13T10:00:00Z",
        }], "nextCursor": None, "hasMore": False}
        vacancy = JobicyAdapter(client).search(["data analyst"], 10)[0]
        self.assertEqual(vacancy.source_id, "20")
        self.assertEqual(vacancy.source_status, "Активна в свежей выдаче")
        self.assertEqual(vacancy.job_type, "full-time, contract")

    def test_arbeitnow_keeps_only_remote(self):
        client = Mock()
        client.get.return_value = {"data": [
            {"slug": "office", "title": "Python Developer", "company_name": "Office",
             "description": "Python", "url": "https://example/office", "remote": False},
            {"slug": "remote", "title": "Python Developer", "company_name": "Remote",
             "description": "Python RAG", "url": "https://example/remote", "remote": True,
             "location": "Europe", "tags": ["Python"], "job_types": ["full-time"]},
        ], "links": {"next": None}}
        vacancies = ArbeitnowAdapter(client).search(["Python"], 10)
        self.assertEqual([item.source_id for item in vacancies], ["remote"])



class SpeechRoutingTests(unittest.TestCase):
    def test_uses_remote_when_available(self):
        remote = Mock()
        remote.is_available.return_value = True
        remote.transcribe.return_value = TranscriptionResult("remote text", "remote", "model")
        local = Mock()
        result = TranscriptionService(remote, local).transcribe(Path("audio.wav"))
        self.assertEqual(result.backend, "remote")
        local.transcribe.assert_not_called()

    def test_falls_back_when_health_check_fails(self):
        remote = Mock()
        remote.is_available.return_value = False
        local = Mock()
        local.transcribe.return_value = TranscriptionResult("local text", "local", "model")
        result = TranscriptionService(remote, local).transcribe(Path("audio.wav"))
        self.assertEqual(result.backend, "local")


class MigrationTests(unittest.TestCase):
    def test_migrates_synthetic_tracker_and_catalog(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job_dir = root / "job-search"
            interview_dir = root / "interview-coach"
            vacancies_dir = job_dir / "vacancies"
            vacancies_dir.mkdir(parents=True)
            interview_dir.mkdir()
            tracker_path = job_dir / "tracker.xlsx"

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Вакансии"
            sheet.append([
                "Компания", "Позиция", "Направление", "Тип", "Формат", "Зарплата",
                "Ссылка", "Статус", "Комментарий", "Дата отклика",
            ])
            sheet.append([
                "Example", "LLM developer", "AI/LLM", "Вакансия", "Удаленно", "",
                "https://hh.ru/vacancy/42", "Подал", "", "2026-10-13",
            ])
            workbook.save(tracker_path)
            (vacancies_dir / "hh_catalog.json").write_text(json.dumps({
                "42": {
                    "source": "hh.ru", "source_id": "42", "employer": "Example",
                    "name": "LLM developer", "url": "https://hh.ru/vacancy/42",
                    "source_status": "Активна", "checked_at": "2026-10-13T12:00:00+03:00",
                    "description": "RAG Python", "assessment": {"score": 80, "direction": "AI/LLM"},
                }
            }, ensure_ascii=False), encoding="utf-8")
            repository = CareerRepository(root / "career.db")
            with patch("career_agent.migration.JOB_DIR", job_dir), \
                    patch("career_agent.migration.INTERVIEW_DIR", interview_dir), \
                    patch("career_agent.migration.TRACKER_PATH", tracker_path), \
                    patch("career_agent.migration.export_tracker"):
                counts = LegacyMigrator(repository).apply()
            rows = repository.list_vacancies()
            self.assertEqual(counts["tracker_rows"], 1)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["application_status"], "Подал")
            self.assertEqual(rows[0]["description"], "RAG Python")


if __name__ == "__main__":
    unittest.main()
