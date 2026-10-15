from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..paths import TRACKER_PATH
from ..repositories.sqlite import CareerRepository


HEADERS = [
    "Компания", "Позиция", "Направление", "Тип", "Формат", "Зарплата",
    "Ссылка", "Статус", "Комментарий", "Дата отклика", "Источник",
    "ID источника", "Статус источника", "Проверено", "Оценка",
]


def export_tracker(repository: CareerRepository, path: Path = TRACKER_PATH) -> Path:
    rows = repository.list_vacancies()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Вакансии"
    for column, name in enumerate(HEADERS, 1):
        cell = sheet.cell(1, column, name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    for row in rows:
        sheet.append([
            row["company"], row["title"], row["direction"], row["job_type"],
            row["work_format"], row["salary"], row["url"],
            row["application_status"] or "Не подал", row["comment"] or "",
            row["applied_at"] or "", row["source"], row["source_id"],
            row["source_status"], row["checked_at"], row["score"],
        ])
    for cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 45)
        sheet.column_dimensions[cells[0].column_letter].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.{datetime.now().timestamp():.0f}.tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    return path
