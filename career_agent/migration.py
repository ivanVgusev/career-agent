import csv
import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .domain.models import Assessment, Vacancy
from .exporters.excel import export_tracker
from .paths import INTERVIEW_DIR, JOB_DIR, TRACKER_PATH
from .repositories.sqlite import CareerRepository


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _stable_legacy_id(company: str, title: str, url: str) -> str:
    value = f"{company}\0{title}\0{url}".encode("utf-8")
    return "legacy-" + hashlib.sha256(value).hexdigest()[:16]


class LegacyMigrator:
    def __init__(self, repository: CareerRepository):
        self.repository = repository

    def inspect(self) -> dict[str, int]:
        counts = {"catalog_vacancies": 0, "tracker_rows": 0, "interview_rows": 0}
        catalog_path = JOB_DIR / "vacancies" / "hh_catalog.json"
        if catalog_path.exists():
            counts["catalog_vacancies"] = len(json.loads(catalog_path.read_text(encoding="utf-8")))
        if TRACKER_PATH.exists():
            workbook = load_workbook(TRACKER_PATH, read_only=True, data_only=True)
            counts["tracker_rows"] = max(0, workbook["Вакансии"].max_row - 1)
            workbook.close()
        for path in (INTERVIEW_DIR / "scoreboard.csv", INTERVIEW_DIR / "speech_metrics.csv"):
            if path.exists():
                with path.open(encoding="utf-8", newline="") as handle:
                    counts["interview_rows"] += sum(1 for _ in csv.DictReader(handle))
        return counts

    def apply(self) -> dict[str, int]:
        self.repository.initialize()
        if self.repository.get_metadata("legacy_migrated_at"):
            raise RuntimeError("Миграция уже выполнена для этой базы")
        counts = {"catalog_vacancies": 0, "tracker_rows": 0, "interview_rows": 0}
        self._migrate_catalog(counts)
        self._migrate_tracker(counts)
        self._migrate_interviews(counts)
        timestamp = _now()
        self.repository.set_metadata("legacy_migrated_at", timestamp)
        if TRACKER_PATH.exists():
            backup = TRACKER_PATH.with_name(
                f"tracker.pre-sqlite-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            shutil.copy2(TRACKER_PATH, backup)
        export_tracker(self.repository)
        return counts

    def _migrate_catalog(self, counts: dict[str, int]) -> None:
        path = JOB_DIR / "vacancies" / "hh_catalog.json"
        if not path.exists():
            return
        for record in json.loads(path.read_text(encoding="utf-8")).values():
            assessment_raw = record.get("assessment") or {}
            assessment = Assessment(
                int(assessment_raw.get("score") or 0),
                assessment_raw.get("direction") or "",
                assessment_raw.get("reasons") or [], assessment_raw.get("warnings") or [],
                bool(assessment_raw.get("rejected")),
            )
            vacancy = Vacancy(
                source=record.get("source") or "hh.ru",
                source_id=str(record.get("source_id")),
                company=record.get("employer") or record.get("company") or "Не указан",
                title=record.get("name") or record.get("title") or "Без названия",
                url=record.get("url") or "", source_status=record.get("source_status") or "Неизвестна",
                checked_at=record.get("checked_at") or _now(), direction=assessment.direction,
                work_format=record.get("schedule") or "", salary=record.get("salary") or "Не указана",
                score=assessment.score, area=record.get("area") or "",
                employment=record.get("employment") or "", experience=record.get("experience") or "",
                published_at=record.get("published_at"), description=record.get("description") or "",
                key_skills=record.get("key_skills") or [], assessment=assessment,
            )
            self.repository.upsert_vacancy(vacancy)
            counts["catalog_vacancies"] += 1

    def _migrate_tracker(self, counts: dict[str, int]) -> None:
        if not TRACKER_PATH.exists():
            return
        workbook = load_workbook(TRACKER_PATH, read_only=True, data_only=True)
        sheet = workbook["Вакансии"]
        headers = {cell.value: cell.column - 1 for cell in sheet[1] if cell.value}

        def value(row: tuple[Any, ...], name: str, default: Any = "") -> Any:
            index = headers.get(name)
            return row[index] if index is not None and index < len(row) and row[index] is not None else default

        for row in sheet.iter_rows(min_row=2, values_only=True):
            company, title, url = str(value(row, "Компания")), str(value(row, "Позиция")), str(value(row, "Ссылка"))
            source = str(value(row, "Источник", "manual"))
            source_id = str(value(row, "ID источника", ""))
            hh_match = re.search(r"hh\.ru/vacancy/(\d+)", url)
            if not source_id and hh_match:
                source, source_id = "hh.ru", hh_match.group(1)
            source_id = source_id or _stable_legacy_id(company, title, url)
            existing = self.repository.get_vacancy(source, source_id)
            existing_value = lambda name, fallback: existing[name] if existing and existing[name] not in (None, "") else fallback
            vacancy = Vacancy(
                source=source, source_id=source_id, company=company, title=title, url=url,
                source_status=str(value(row, "Статус источника", existing_value("source_status", "Неизвестна"))),
                checked_at=str(value(row, "Проверено", existing_value("checked_at", _now()))),
                direction=str(value(row, "Направление", existing_value("direction", ""))),
                job_type=str(value(row, "Тип", existing_value("job_type", "Вакансия"))),
                work_format=str(value(row, "Формат", existing_value("work_format", ""))),
                salary=str(value(row, "Зарплата", existing_value("salary", "Не указана"))),
                score=int(value(row, "Оценка", existing_value("score", 0)) or 0),
                description=existing_value("description", ""),
            )
            self.repository.upsert_vacancy(vacancy)
            status = str(value(row, "Статус", "Не подал"))
            comment = str(value(row, "Комментарий", ""))
            applied_at = value(row, "Дата отклика", None)
            self.repository.set_application_status(
                source, source_id, status, comment,
                applied_at=str(applied_at) if applied_at else None,
            )
            counts["tracker_rows"] += 1
        workbook.close()

    def _migrate_interviews(self, counts: dict[str, int]) -> None:
        for path in (INTERVIEW_DIR / "scoreboard.csv", INTERVIEW_DIR / "speech_metrics.csv"):
            if not path.exists():
                continue
            with path.open(encoding="utf-8", newline="") as handle:
                for index, row in enumerate(csv.DictReader(handle), 1):
                    key = row.get("session") or row.get("name") or f"{path.stem}-{index}"
                    payload = {"source_file": path.name, **row}
                    self.repository.upsert_interview_payload(key, payload)
                    counts["interview_rows"] += 1
